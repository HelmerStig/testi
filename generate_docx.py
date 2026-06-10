#!/usr/bin/env python3
import os
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Directory con i file txt
testi_dir = Path(__file__).parent

# Leggi tutti i file .txt (escluso concept.txt che contiene tutti i testi)
txt_files = sorted([f for f in testi_dir.glob("*.txt") if f.name != "concept.txt"])

for txt_file in txt_files:
    print(f"Generando {txt_file.stem}.docx...")
    
    # Leggi il contenuto del file txt
    with open(txt_file, "r", encoding="utf-8") as f:
        content = f.read()
    
    # Crea un nuovo documento Word
    doc = Document()
    
    # Aggiungi il contenuto
    lines = content.split("\n")
    
    for i, line in enumerate(lines):
        if i == 0:  # Titolo
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line)
            run.font.size = Pt(18)
            run.font.bold = True
        elif line.strip() == "":  # Linea vuota
            doc.add_paragraph()
        else:
            p = doc.add_paragraph(line)
            p.paragraph_format.line_spacing = 1.5
    
    # Salva il documento
    docx_path = txt_file.with_suffix(".docx")
    doc.save(docx_path)
    print(f"✓ {docx_path.name} creato con successo")

print(f"\nTotale: {len(txt_files)} file generati")
